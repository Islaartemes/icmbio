from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from pg8000.dbapi import Connection


NOME_PLANILHA = "vw_4_divulgacao_embargos_eletro"
TAMANHO_LOTE = 500

COLUNAS_FONTE = (
    "vw_num_emb",
    "numero_emb",
    "serie",
    "origem",
    "numero_ai",
    "cpf_cnpj",
    "autuado",
    "desc_infra",
    "desc_sanc",
    "artigo_1",
    "artigo_2",
    "tipo_infra",
    "nome_uc",
    "cnuc",
    "municipio",
    "uf",
    "data",
    "ano",
    "obs",
    "julgamento",
    "area",
    "processo",
)


class ErroEstruturaArquivo(RuntimeError):
    pass


def _converter_para_texto(valor: object) -> str | None:
    if valor is None:
        return None

    if isinstance(valor, str):
        if not valor.strip():
            return None
        return valor

    if isinstance(valor, datetime):
        return valor.isoformat(sep=" ")

    if isinstance(valor, date):
        return valor.isoformat()

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor)


def _validar_cabecalho(cabecalho: tuple[object, ...]) -> None:
    colunas_encontradas = tuple(
        str(valor).strip() if valor is not None else ""
        for valor in cabecalho
    )

    if colunas_encontradas != COLUNAS_FONTE:
        raise ErroEstruturaArquivo(
            "As colunas do XLSX sao diferentes das esperadas. "
            f"Encontradas: {colunas_encontradas}"
        )


def carregar_raw(
    conexao: Connection,
    caminho_arquivo: Path,
    id_carga: UUID,
    arquivo_origem: str,
) -> int:
    comando = """
        INSERT INTO raw.embargo_icmbio (
            id_carga,
            arquivo_origem,
            linha_origem,
            vw_num_emb,
            numero_emb,
            serie,
            origem,
            numero_ai,
            cpf_cnpj,
            autuado,
            desc_infra,
            desc_sanc,
            artigo_1,
            artigo_2,
            tipo_infra,
            nome_uc,
            cnuc,
            municipio,
            uf,
            data,
            ano,
            obs,
            julgamento,
            area,
            processo
        )
        VALUES (
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s
        )
    """

    arquivo_excel = load_workbook(
        caminho_arquivo,
        read_only=True,
        data_only=True,
    )

    try:
        if NOME_PLANILHA not in arquivo_excel.sheetnames:
            raise ErroEstruturaArquivo(
                f"Planilha nao encontrada: {NOME_PLANILHA}"
            )

        planilha = arquivo_excel[NOME_PLANILHA]
        linhas = planilha.iter_rows(values_only=True)
        cabecalho = next(linhas, None)

        if cabecalho is None:
            raise ErroEstruturaArquivo("A planilha esta vazia")

        _validar_cabecalho(cabecalho)

        cursor = conexao.cursor()
        quantidade_linhas = 0
        lote: list[tuple[object, ...]] = []

        try:
            for numero_linha, linha in enumerate(linhas, start=2):
                valores = list(linha[: len(COLUNAS_FONTE)])
                valores += [None] * (
                    len(COLUNAS_FONTE) - len(valores)
                )

                if all(
                    valor is None
                    or (isinstance(valor, str) and not valor.strip())
                    for valor in valores
                ):
                    continue

                valores_texto = tuple(
                    _converter_para_texto(valor)
                    for valor in valores
                )

                lote.append(
                    (
                        id_carga,
                        arquivo_origem,
                        numero_linha,
                        *valores_texto,
                    )
                )

                if len(lote) == TAMANHO_LOTE:
                    cursor.executemany(comando, lote)
                    quantidade_linhas += len(lote)
                    lote.clear()

            if lote:
                cursor.executemany(comando, lote)
                quantidade_linhas += len(lote)

            return quantidade_linhas
        finally:
            cursor.close()
    finally:
        arquivo_excel.close()